import streamlit as st
from dotenv import load_dotenv
from PyPDF2 import PdfReader
from langchain.text_splitter import CharacterTextSplitter
from langchain.embeddings import OpenAIEmbeddings
from langchain.vectorstores import FAISS
from langchain.chat_models import ChatOpenAI
from langchain.memory import ConversationBufferMemory
from langchain.chains import ConversationalRetrievalChain
from htmlTemplates import css, bot_template, user_template
from googleapiclient.discovery import build
from google_auth_oauthlib.flow import InstalledAppFlow
import os
import json
from google.auth.transport.requests import Request
import pickle
from googleapiclient.http import MediaIoBaseDownload
import io
import docx
import pytesseract
from PIL import Image
import easyocr
import openpyxl
import requests
from github import Github
from streamlit_chat import message
#scopes for Google Drive API
SCOPES = ['https://www.googleapis.com/auth/drive']
my_dict = []
def get_authenticated_service():
    """Authenticate and create a Google Drive API service."""
    creds = None

    # Check if we already have valid credentials in a pickle file
    if os.path.exists('token.pickle'):
        with open('token.pickle', 'rb') as token_file:
            creds = pickle.load(token_file)

    # If there are no valid credentials, authenticate the user interactively
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            g = Github('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')
            repo = g.get_repo('araj8899/gdrive')
            contents = repo.get_contents('token.json')
            decoded = contents.decoded_content
            with open("token.json", "wb") as f:
                    f.write(decoded)
            flow = InstalledAppFlow.from_client_secrets_file('token.json', SCOPES)
            creds = flow.run_local_server(port=0)
        
        # Save the credentials for future use
        with open('token.pickle', 'wb') as token_file:
            pickle.dump(creds, token_file)

    # Create and return the Drive API service
    return build('drive', 'v3', credentials=creds)

def download_files(service, file_id, output_path):
    request = service.files().get_media(fileId=file_id)
    with io.FileIO(output_path, mode='wb') as file:
        downloader = MediaIoBaseDownload(file, request)
        done = False
        while done is False:
            status, done = downloader.next_chunk()
def read_pdf_content(pdf_path):
    text = ""
    with open(pdf_path, "rb") as file:
        pdf_reader = PdfReader(file)
        for page in pdf_reader.pages:
            text += page.extract_text()
    return text

def read_docx(file_path):
    doc = docx.Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def read_image(image_file_name):
    # myconfig = r"--psm 11 --oem 3"
    # pytesseract.pytesseract.tesseract_cmd = 'C:/Program Files/Tesseract-OCR/tesseract.exe'
    # image = Image.open(image_file_name)
    # image_content = pytesseract.image_to_string(image, config= myconfig)
    # return image_content
    reader = easyocr.Reader(['en'])
    result = reader.readtext(image_file_name)
    extracted_text = [text for (_, text, _) in result]
    extracted_text_string = " ".join(extracted_text)
    return extracted_text_string

# def read_excel(excel_file_name): 
#     workbook = openpyxl.load_workbook(excel_file_name)
#     worksheet = workbook["Mayurs KT"]
#     extracted_text = []
#     for row in worksheet.iter_rows(values_only=True):
#         for cell_value in row:
#             if cell_value:
#                 extracted_text.append(cell_value)
#     for text in extracted_text:
#         print(text)

def read_file_contents(file_id):
    """Read the contents of a file given its file ID."""
    service = get_authenticated_service()

    try:
        # Get the file content
        file_content = service.files().get_media(fileId=file_id).execute()

        # Decode the content and return as a string
        content_str = file_content.decode('utf-8')
        return content_str
    except Exception as e:
        print(f"An error occurred: {e}")
        return None

def get_file_link(drive_service,file_id):
   url = drive_service.files().get(fileId=file_id, fields="webViewLink").execute()
   return url.get("webViewLink")

def fetch_files_with_format(drive_service, mime_type):
    results = drive_service.files().list(q=f"mimeType='{mime_type}'", fields="files(id, name)").execute()
    items = results.get('files', [])
    return items
def get_pdf_text(pdf_docs):
    text = ""
    for pdf in pdf_docs:
        pdf_reader = PdfReader(pdf)
        for page in pdf_reader.pages:
            text += page.extract_text()
    return text


def get_text_chunks(text):
    text_splitter = CharacterTextSplitter(
        separator="\n",
        chunk_size=1000,
        chunk_overlap=200,
        length_function=len
    )
    chunks = text_splitter.split_text(text)
    return chunks


def get_vectorstore(text_chunks):
    embeddings = OpenAIEmbeddings()
    vectorstore = FAISS.from_texts(texts=text_chunks, embedding=embeddings)
    return vectorstore


def get_conversation_chain(vectorstore):
    llm = ChatOpenAI()

    memory = ConversationBufferMemory(
        memory_key='chat_history', return_messages=True)
    conversation_chain = ConversationalRetrievalChain.from_llm(
        llm=llm,
        retriever=vectorstore.as_retriever(),
        memory=memory
    )
    return conversation_chain


def handle_userinput(user_question):
    response = st.session_state.conversation({'question': user_question})
    st.session_state.chat_history = response['chat_history']

    for i, msg in enumerate(st.session_state.chat_history):
        if i % 2 == 0:
            message(msg.content, is_user=True)
        else:
            message(msg.content, is_user=False)

def gdrive_api(): 
   service = get_authenticated_service()
   pdf_files = fetch_files_with_format(service, 'application/pdf')
   for file in pdf_files:
       print(f"{file['name']} - {file['id']}")
       
       pdf_file_id = file['id']
       pdf_file_name = file['name']

       download_files(service, pdf_file_id, pdf_file_name)
       pdf_content = read_pdf_content(pdf_file_name)
       pdf_content_added = pdf_file_name + " file contents : "+ pdf_content 
       my_dict.append(pdf_content_added)

#    plain_files = fetch_files_with_format(service, 'text/plain')
#    for file in plain_files:
#        print(f"{file['name']} - {file['id']}")
#        plain_file_name = file['name']
#        plain_file_id = file['id']
#        plain_file_content = read_file_contents(plain_file_id)
#        plain_content_added = plain_file_name + " file contents : "+ plain_file_content
#        my_dict.append(plain_content_added) 

#    doc_files = fetch_files_with_format(service, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
#    for file in doc_files:
#        print(f"{file['name']} - {file['id']}")
#        doc_file_id = file['id']
#        doc_file_name = file['name']
#        download_files(service, doc_file_id, doc_file_name)
#        doc_content = read_docx(doc_file_name)
#        doc_content_added = doc_file_name + " file contents : "+ doc_content
#        my_dict.append(doc_content_added)
    
   
#    image_png_files = fetch_files_with_format(service, 'image/png')
#    for file in image_png_files:
#        print(f"{file['name']} - {file['id']}")
#        image_file_id = file['id']
#        image_file_name = file['name']
#        download_files(service, image_file_id, image_file_name)
#        image_content = read_image(image_file_name)
#        image_content_added = image_file_name + " file contents : "+ image_content
#        my_dict.append(image_content_added)
    
#    image_jpeg_files = fetch_files_with_format(service, 'image/jpeg')
#    for file in image_jpeg_files:
#       print(f"{file['name']} - {file['id']}")
#       image_file_id = file['id']
#       image_file_name = file['name']
#       download_files(service, image_file_id, image_file_name)
#       image_content = read_image(image_file_name)
#       image_content_added = image_file_name + " file contents : "+ image_content
#       my_dict.append(image_content_added)
    
#    excel_files = fetch_files_with_format(service, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
#    for file in excel_files:
#        print(f"{file['name']} - {file['id']}")
#        excel_file_id = file['id']
#        excel_file_name = file['name']
#        download_files(service, excel_file_id, excel_file_name)
#        excel_content = read_excel(excel_file_name)
#        excel_content_added = excel_file_name + " file contents : "+ excel_content
#        my_dict.append(excel_content_added)

def main():
    load_dotenv()
    st.set_page_config(page_title="PhantomAI",
                       page_icon=":books:")
    
    st.write(css, unsafe_allow_html=True)
    if "conversation" not in st.session_state:
        st.session_state.conversation = None
    if "chat_history" not in st.session_state:
        st.session_state.chat_history = None
    st.header("Chat with ChatBotAi :book:")
    with st.sidebar: 
       user_question = st.text_input("Ask me a question :")
    if user_question:
        handle_userinput(user_question)
    if st.button("START"):
      gdrive_api()
      my_files_str = json.dumps(my_dict)
      print(my_files_str)
      with st.spinner("Thinking...."):

          # get the text chunks
          text_chunks = get_text_chunks(my_files_str)

          # create vector store
          vectorstore = get_vectorstore(text_chunks)

          # create conversation chain
          st.session_state.conversation = get_conversation_chain(
              vectorstore)


if __name__ == '__main__':
    main()
